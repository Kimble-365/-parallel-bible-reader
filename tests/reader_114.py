"""Real new-language reading, note migration, Office downloads, audio scope and offline use."""
from pathlib import Path
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from threading import Thread
from zipfile import ZipFile
from docx import Document
from pptx import Presentation
from playwright.sync_api import sync_playwright
import json

ROOT=Path(__file__).resolve().parents[1];OUT=Path('/tmp/reader-browser-office-114');OUT.mkdir(exist_ok=True)
class H(SimpleHTTPRequestHandler):
    def __init__(self,*a,**k):super().__init__(*a,directory=str(ROOT),**k)
    def log_message(self,*a):pass
server=ThreadingHTTPServer(('127.0.0.1',0),H);Thread(target=server.serve_forever,daemon=True).start();base=f'http://127.0.0.1:{server.server_port}'
ids=['jpn-bungo','mya-judson','vie-1934','ceb-pinadayag'];url=base+'/releases/v1.14/?edition=jpn-bungo&compare1=mya-judson&compare2=vie-1934&compare3=ceb-pinadayag#JHN.3.16'
def ready(page):
    page.wait_for_function("typeof document.querySelector('#notes-open').onclick==='function' && document.querySelector('#s1-v16 .main-text')")
def download(page,id,name):
    with page.expect_download() as event:page.click('#'+id)
    path=OUT/name;event.value.save_as(path);return path
try:
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox']);ctx=browser.new_context(accept_downloads=True);page=ctx.new_page();errors=[];requests=[];page.on('pageerror',lambda e:errors.append(str(e)));page.on('request',lambda r:requests.append(r.url))
        # Existing 1.13 notes must survive upgrading within the same browser/origin.
        page.goto(base+'/releases/v1.13/?edition=engwebp&compare1=&compare2=&compare3=#JHN.3.16');page.wait_for_function("typeof document.querySelector('#notes-open').onclick==='function'");page.click('#notes-open');page.wait_for_selector('.notes-quote');page.fill('#notes-body','Saved in reader 1.13')
        page.goto(url);ready(page);assert page.evaluate('BibleCatalog.editions.length')==55;assert page.evaluate('BibleCatalog.audio_editions.length')==34;assert page.locator('.source-pane').count()==4;assert 'only kana' in page.locator('#text-form-notice').inner_text()
        source=page.evaluate("async()=>{const b=await ReaderOffline.loadBook('JHN',['jpn-bungo']);return b.editions['jpn-bungo'].find(r=>r[0]===3&&r[1]===16)[2]}")
        page.click('#notes-open');assert page.locator('#notes-body').input_value()=='Saved in reader 1.13';assert 'WEB' in page.locator('.notes-quote').inner_text()
        page.click('#notes-new');page.wait_for_function("document.querySelector('.notes-quote')?.textContent.includes('Japanese')");assert source in page.locator('.notes-quote').inner_text();page.fill('#notes-title','Japanese study & <literal>');page.fill('#notes-body','Own writing äöå 日本語 Ελληνικά אֱלֹהִים')
        for id,name in [('notes-export-text','current.txt'),('notes-export-docx','current.docx'),('notes-export-pptx','current.pptx')]:download(page,id,name)
        assert source in (OUT/'current.txt').read_text(encoding='utf-8');assert source in '\n'.join(x.text for x in Document(OUT/'current.docx').paragraphs)
        deck=Presentation(OUT/'current.pptx');assert source in '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        for name in ['current.docx','current.pptx']:
            with ZipFile(OUT/name) as z:assert z.testzip() is None
        assert not [r for r in requests if r.startswith(('https://docs.google','https://slides.google','https://www.jesusfilm.org'))], 'Exports must not upload private notes or preload a film'
        page.get_by_text('Open in Google Docs or Google Slides',exact=True).click();assert page.locator('#notes-google-help a[href="https://docs.google.com/document/"]').count()==1;assert page.locator('#notes-google-help a[href="https://slides.google.com/"]').count()==1
        page.click('#notes-close');page.locator('#film-panel').evaluate('(e)=>e.open=true');page.select_option('#film-language','eng');page.click('#film-match');assert page.locator('#film-language').input_value()=='jpn';assert page.locator('#film-watch').get_attribute('href').endswith('/japanese.html');assert page.locator('#film-watch').get_attribute('target')=='_blank';page.select_option('#film-language','ceb');assert 'cebuano.html' in page.locator('#film-watch').get_attribute('href')
        page.locator('#provider-panel').evaluate('(e)=>e.open=true');assert page.locator('#provider-links a[href="https://www.bible.com/languages/jpn"]').count()==1;assert page.locator('#provider-links a[href*="JPNJCB"]').count()==1
        # An OT-only audio edition must never be redirected to Matthew.
        page.locator('#audio-panel').evaluate('(e)=>e.open=true');page.select_option('#audio-edition','MYAOJV');assert page.locator('#audio-external').get_attribute('href').endswith('/MYAOJV/GEN/1');assert 'Old Testament' in page.locator('#audio-location').inner_text();page.select_option('#audio-edition','JPNJCB');assert page.locator('#audio-external').get_attribute('href').endswith('/JPNJCB/JHN/3')
        page.locator('#appearance-settings').evaluate('(e)=>e.open=true');page.select_option('#text-display','original');page.wait_for_function("document.querySelector('#text-form-notice').textContent===''");page.locator('#search-panel').evaluate('(e)=>e.open=true');page.select_option('#scope','chapter');page.fill('#search','Đức');page.click('#search-button');page.wait_for_function("document.querySelector('#results > p')?.textContent.startsWith('Complete')");assert page.locator('#results .result[data-edition="vie-1934"][data-verse="16"]').count()==1;page.check('#search-exact-marks');page.click('#search-button');page.wait_for_function("document.querySelector('#results > p')?.textContent.startsWith('Complete')");assert page.locator('#results .result[data-edition="vie-1934"][data-verse="16"]').count()==0
        page.select_option('#primary','mya-judson');page.wait_for_function("document.querySelector('#s1-v16 [data-edition=\"mya-judson\"]')");page.click('#notes-open');page.click('#notes-new');page.wait_for_function("document.querySelector('#notes-title').value==='John 3:16' && document.querySelector('.notes-quote')?.textContent.includes('Burmese')");page.fill('#notes-body','A second language note');page.select_option('#notes-export-scope','all');path=download(page,'notes-export-docx','all.docx');alltext='\n'.join(p.text for p in Document(path).paragraphs);assert 'Saved in reader 1.13' in alltext and 'Japanese study' in alltext and 'Burmese' in alltext
        page.evaluate("async ids=>{await ReaderOffline.download(ids)}",ids);saved=page.evaluate('ReaderOffline.list()');assert len([r for r in saved if r['id'] in ids])==264
        page.evaluate('navigator.serviceWorker.ready');page.wait_for_function('!!navigator.serviceWorker.controller');ctx.set_offline(True);page.reload();ready(page);page.click('#notes-open');page.select_option('#notes-list',index=2);assert 'A second language note' in page.locator('#notes-body').input_value();page.select_option('#notes-export-scope','all');offline=download(page,'notes-export-pptx','offline-all.pptx');assert len(Presentation(offline).slides)>=6;assert page.locator('#film-language').input_value()=='ceb'
        page.set_viewport_size({'width':390,'height':844});assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+1');assert not errors,errors;browser.close()
        print('PASS: 55 editions / 34 audio choices; real four-language source panes; partial Japanese labels; 1.13 notes persist; Text 1 quotes; actual browser UTF-8/DOCX/PPTX exports independently open; Google manual-import links; no automatic uploads/film load; official film language links; provider links; OT-only audio fallback; Vietnamese historical-letter matching; all four complete Bibles saved (264 books); offline Office exports; mobile overflow; no page errors')
finally:server.shutdown()
