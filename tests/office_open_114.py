"""Read browser-produced Office files with independent document libraries."""
from pathlib import Path, PurePosixPath
from zipfile import ZipFile
import json, posixpath, xml.etree.ElementTree as E
from docx import Document
from pptx import Presentation

base=Path('/tmp/reader-office-114');expected=json.loads((base/'expected.json').read_text());q,n,another=(expected[k] for k in ['q','n','another'])
for filename in ['notes.docx','notes.pptx']:
    with ZipFile(base/filename) as z:
        assert z.testzip() is None
        names=set(z.namelist())
        types=z.read('[Content_Types].xml').decode();assert '<Types xmlns=' in types and types.startswith('<?xml'), 'Use the Office-compatible default content-types namespace'
        for name in names:
            if name.endswith(('.xml','.rels')):
                root=E.fromstring(z.read(name))
                if name.endswith('.rels'):
                    directory=posixpath.dirname(posixpath.dirname(name)) if '/_rels/' in name else ''
                    for rel in root:
                        if rel.get('TargetMode')!='External':
                            target=posixpath.normpath(posixpath.join(directory,rel.get('Target')))
                            assert target in names,(name,target)
doc=Document(base/'notes.docx');paragraphs=[p.text for p in doc.paragraphs];text='\n'.join(paragraphs)
for value in [n['title'],q['reference']+' — '+q['edition'],q['text'],n['body'],another['title'],another['body']]:assert value in text,value[:50]
assert doc.styles['Quotation'] and doc.styles['Title']
prs=Presentation(base/'notes.pptx');assert len(prs.slides)>4;assert prs.slide_width==12192000 and prs.slide_height==6858000
slide_text='\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
for value in [q['reference'],q['edition'],q['text'],n['title'],another['title'],another['body']]:assert value in slide_text,value[:50]
plain=(base/'notes.txt').read_text(encoding='utf-8');assert q['text'] in plain and n['body'] in plain
print('PASS: DOCX independently opens in python-docx; PPTX independently opens in python-pptx with valid master/layout/slide relationships; all XML/ZIP checks pass; citations, RTL and mixed Unicode retained; widescreen deck; UTF-8 text round-trips')
